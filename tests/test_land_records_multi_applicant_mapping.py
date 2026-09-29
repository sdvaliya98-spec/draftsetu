import os
import sys
import unittest
from pathlib import Path
from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from backend.services.template_service import template_service
from backend.services.docx_engine import (
    render_docx_template,
    _normalize_context,
    convert_docx_to_pdf,
    PREVIEW_WATERMARK_TEXT
)

class TestLandRecordsMultiApplicantMapping(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.template_file = Path(template_service.get_full_path('a7a287b25ae14d00b8de1de9ce9afd90.docx'))
        cls.output_dir = PROJECT_ROOT / "scratch" / "test_mapping_regressions"
        cls.output_dir.mkdir(parents=True, exist_ok=True)
        assert cls.template_file.exists(), f"NA Template file not found: {cls.template_file}"

        cls.applicants_6 = [
            {"index": "1", "name": "અરજદાર એ (Applicant A)", "address": "અમદાવાદ", "age": "40", "aadhaar": "111122223333"},
            {"index": "2", "name": "અરજદાર બી (Applicant B)", "address": "અમદાવાદ", "age": "42", "aadhaar": "222233334444"},
            {"index": "3", "name": "અરજદાર સી (Applicant C)", "address": "અમદાવાદ", "age": "38", "aadhaar": "333344445555"},
            {"index": "4", "name": "અરજદાર ડી (Applicant D)", "address": "અમદાવાદ", "age": "45", "aadhaar": "444455556666"},
            {"index": "5", "name": "અરજદાર ઈ (Applicant E)", "address": "અમદાવાદ", "age": "50", "aadhaar": "555566667777"},
            {"index": "6", "name": "અરજદાર એફ (Applicant F)", "address": "અમદાવાદ", "age": "35", "aadhaar": "666677778888"},
        ]

        cls.base_data = {
            "VILLAGE_NAME": "બોપલ",
            "TALUKA": "દસ્ક્રોઈ",
            "DISTRICT": "અમદાવાદ",
            "SURVEY_NO": "૧૨૩/૧",
            "LAND_AREA": "૫૦૦",
            "APPLICATION_DATE": "૨૮/૦૯/૨૦૨૬",
            "PLACE": "અમદાવાદ",
            "AFFIDAVIT_DATE": "૨૮/૦૯/૨૦૨૬",
            "APPLICANTS": cls.applicants_6,
            "EXTRA_PARAGRAPHS": [{"text": "અમો ખેડૂત ખાતેદાર છીએ."}]
        }

    def test_scenario_1_exact_requirement_indices_1_2_4_6(self):
        """Test exact requirement: 6 applicants, 1 land record, owner_applicant_indices = [1, 2, 4, 6]"""
        data = dict(self.base_data)
        data["LAND_RECORDS"] = [
            {
                "index": "1",
                "note_no": "123",
                "note_date": "01/01/2025",
                "note_details": "વારસાઈથી દાખલ",
                "owner_applicant_indices": [1, 2, 4, 6]
            }
        ]
        out_docx = self.output_dir / "reg_scenario_1.docx"
        render_docx_template(str(self.template_file), data, str(out_docx))
        
        doc = Document(str(out_docx))
        table = doc.tables[1]
        self.assertEqual(len(table.rows), 2, "Must produce exactly 1 data row (no row duplication)")
        
        owner_cell = table.rows[1].cells[5].text.strip()
        self.assertIn("Applicant A", owner_cell)
        self.assertIn("Applicant B", owner_cell)
        self.assertIn("Applicant D", owner_cell)
        self.assertIn("Applicant F", owner_cell)
        self.assertNotIn("Applicant C", owner_cell)
        self.assertNotIn("Applicant E", owner_cell)

    def test_scenario_2_indices_1_2(self):
        """Test subset indices [1, 2]"""
        data = dict(self.base_data)
        data["LAND_RECORDS"] = [
            {
                "index": "1",
                "note_no": "456",
                "note_date": "10/05/2020",
                "note_details": "વેચાણ",
                "owner_applicant_indices": [1, 2]
            }
        ]
        out_docx = self.output_dir / "reg_scenario_2.docx"
        render_docx_template(str(self.template_file), data, str(out_docx))
        doc = Document(str(out_docx))
        owner_cell = doc.tables[1].rows[1].cells[5].text.strip()
        self.assertIn("Applicant A", owner_cell)
        self.assertIn("Applicant B", owner_cell)
        self.assertNotIn("Applicant C", owner_cell)

    def test_scenario_3_indices_3_5_6(self):
        """Test subset indices [3, 5, 6]"""
        data = dict(self.base_data)
        data["LAND_RECORDS"] = [
            {
                "index": "1",
                "note_no": "789",
                "note_date": "15/08/2021",
                "note_details": "વારસાઈ",
                "owner_applicant_indices": [3, 5, 6]
            }
        ]
        out_docx = self.output_dir / "reg_scenario_3.docx"
        render_docx_template(str(self.template_file), data, str(out_docx))
        doc = Document(str(out_docx))
        owner_cell = doc.tables[1].rows[1].cells[5].text.strip()
        self.assertIn("Applicant C", owner_cell)
        self.assertIn("Applicant E", owner_cell)
        self.assertIn("Applicant F", owner_cell)
        self.assertNotIn("Applicant A", owner_cell)

    def test_scenario_4_all_applicants_1_to_6(self):
        """Test all 6 applicants selected [1, 2, 3, 4, 5, 6]"""
        data = dict(self.base_data)
        data["LAND_RECORDS"] = [
            {
                "index": "1",
                "note_no": "999",
                "note_date": "01/01/2024",
                "note_details": "સહભાગી",
                "owner_applicant_indices": [1, 2, 3, 4, 5, 6]
            }
        ]
        out_docx = self.output_dir / "reg_scenario_4.docx"
        render_docx_template(str(self.template_file), data, str(out_docx))
        doc = Document(str(out_docx))
        owner_cell = doc.tables[1].rows[1].cells[5].text.strip()
        for letter in ["A", "B", "C", "D", "E", "F"]:
            self.assertIn(f"Applicant {letter}", owner_cell)

    def test_scenario_5_multiple_rows_distinct_mappings(self):
        """Test multiple LAND_RECORD rows with distinct applicant mappings"""
        data = dict(self.base_data)
        data["LAND_RECORDS"] = [
            {"index": "1", "note_no": "101", "note_date": "01/01/2020", "note_details": "નોંધ ૧", "owner_applicant_indices": [1, 2]},
            {"index": "2", "note_no": "202", "note_date": "02/02/2022", "note_details": "નોંધ ૨", "owner_applicant_indices": [4, 6]},
            {"index": "3", "note_no": "303", "note_date": "03/03/2024", "note_details": "નોંધ ૩", "owner_applicant_indices": [3, 5]}
        ]
        out_docx = self.output_dir / "reg_scenario_5.docx"
        render_docx_template(str(self.template_file), data, str(out_docx))
        doc = Document(str(out_docx))
        table = doc.tables[1]
        self.assertEqual(len(table.rows), 4) # 1 header + 3 records
        
        row1_owner = table.rows[1].cells[5].text.strip()
        row2_owner = table.rows[2].cells[5].text.strip()
        row3_owner = table.rows[3].cells[5].text.strip()
        
        self.assertIn("Applicant A", row1_owner)
        self.assertIn("Applicant B", row1_owner)
        self.assertIn("Applicant D", row2_owner)
        self.assertIn("Applicant F", row2_owner)
        self.assertIn("Applicant C", row3_owner)
        self.assertIn("Applicant E", row3_owner)

    def test_scenario_6_backward_compatibility_legacy_owner_name(self):
        """Test legacy data with owner_name string and no owner_applicant_indices"""
        data = dict(self.base_data)
        data["LAND_RECORDS"] = [
            {
                "index": "1",
                "note_no": "888",
                "note_date": "01/01/2018",
                "note_details": "જૂની વારસાઈ નોંધ",
                "owner_name": "શ્રી રમેશભાઈ ગોવિંદભાઈ પટેલ"
            }
        ]
        out_docx = self.output_dir / "reg_scenario_6.docx"
        render_docx_template(str(self.template_file), data, str(out_docx))
        doc = Document(str(out_docx))
        owner_cell = doc.tables[1].rows[1].cells[5].text.strip()
        self.assertIn("શ્રી રમેશભાઈ ગોવિંદભાઈ પટેલ", owner_cell)

    def test_scenario_7_invalid_indices_resilience(self):
        """Test that invalid, out-of-bounds, or string indices are handled safely without crashing"""
        data = dict(self.base_data)
        data["LAND_RECORDS"] = [
            {
                "index": "1",
                "note_no": "555",
                "note_date": "01/01/2023",
                "note_details": "ટેસ્ટ નોંધ",
                "owner_applicant_indices": [1, 99, -2, "invalid", 4],
                "owner_name": "Fallback Name"
            }
        ]
        out_docx = self.output_dir / "reg_scenario_7.docx"
        render_docx_template(str(self.template_file), data, str(out_docx))
        doc = Document(str(out_docx))
        owner_cell = doc.tables[1].rows[1].cells[5].text.strip()
        self.assertIn("Applicant A", owner_cell)
        self.assertIn("Applicant D", owner_cell)
        self.assertNotIn("Fallback Name", owner_cell)

    def test_scenario_8_dynamic_count_10_applicants(self):
        """Test dynamic scaling with 10 applicants (not hardcoded to 6)"""
        applicants_10 = [
            {"index": str(i + 1), "name": f"Applicant_{i + 1}", "address": "City", "age": "30", "aadhaar": f"1000{i}"}
            for i in range(10)
        ]
        data = dict(self.base_data)
        data["APPLICANTS"] = applicants_10
        data["LAND_RECORDS"] = [
            {
                "index": "1",
                "note_no": "1001",
                "note_date": "01/01/2025",
                "note_details": "ટેસ્ટ ૧૦",
                "owner_applicant_indices": [1, 5, 8, 10]
            }
        ]
        out_docx = self.output_dir / "reg_scenario_8.docx"
        render_docx_template(str(self.template_file), data, str(out_docx))
        doc = Document(str(out_docx))
        owner_cell = doc.tables[1].rows[1].cells[5].text.strip()
        self.assertIn("Applicant_1", owner_cell)
        self.assertIn("Applicant_5", owner_cell)
        self.assertIn("Applicant_8", owner_cell)
        self.assertIn("Applicant_10", owner_cell)
        self.assertNotIn("Applicant_2", owner_cell)
        self.assertNotIn("Applicant_9", owner_cell)

    def test_scenario_9_manual_names_only(self):
        """Test manual occupant names only without any applicant index selected"""
        data = dict(self.base_data)
        data["LAND_RECORDS"] = [
            {
                "index": "1",
                "note_no": "456",
                "note_date": "10/05/2020",
                "note_details": "અન્ય ખાતેદાર",
                "owner_applicant_indices": [],
                "owner_other_names": ["XYZ વ્યક્તિ", "ABC વ્યક્તિ"]
            }
        ]
        out_docx = self.output_dir / "reg_scenario_9.docx"
        render_docx_template(str(self.template_file), data, str(out_docx))
        doc = Document(str(out_docx))
        table = doc.tables[1]
        self.assertEqual(len(table.rows), 2, "Must produce exactly 1 data row")
        owner_cell = table.rows[1].cells[5].text.strip()
        self.assertEqual(owner_cell, "XYZ વ્યક્તિ, ABC વ્યક્તિ")

    def test_scenario_10_applicants_plus_manual_names(self):
        """Test both selected applicants [1, 2, 4, 6] AND custom manual names ['XYZ વ્યક્તિ', 'ABC વ્યક્તિ']"""
        data = dict(self.base_data)
        data["LAND_RECORDS"] = [
            {
                "index": "1",
                "note_no": "456",
                "note_date": "10/05/2020",
                "note_details": "સંયુક્ત કબજેદાર",
                "owner_applicant_indices": [1, 2, 4, 6],
                "owner_other_names": ["XYZ વ્યક્તિ", "ABC વ્યક્તિ"]
            }
        ]
        out_docx = self.output_dir / "reg_scenario_10.docx"
        render_docx_template(str(self.template_file), data, str(out_docx))
        doc = Document(str(out_docx))
        table = doc.tables[1]
        self.assertEqual(len(table.rows), 2, "Must produce exactly 1 data row")
        owner_cell = table.rows[1].cells[5].text.strip()
        
        # Verify order: 1. Applicants (A, B, D, F) followed by 2. Other names (XYZ, ABC)
        self.assertIn("Applicant A", owner_cell)
        self.assertIn("Applicant B", owner_cell)
        self.assertIn("Applicant D", owner_cell)
        self.assertIn("Applicant F", owner_cell)
        self.assertIn("XYZ વ્યક્તિ", owner_cell)
        self.assertIn("ABC વ્યક્તિ", owner_cell)
        self.assertNotIn("Applicant C", owner_cell)
        self.assertNotIn("Applicant E", owner_cell)

        # Exact position check: A appears before XYZ
        pos_a = owner_cell.find("Applicant A")
        pos_f = owner_cell.find("Applicant F")
        pos_xyz = owner_cell.find("XYZ વ્યક્તિ")
        pos_abc = owner_cell.find("ABC વ્યક્તિ")
        self.assertTrue(pos_a < pos_f < pos_xyz < pos_abc)

    def test_scenario_11_duplicate_and_whitespace_manual_names(self):
        """Test handling of duplicate manual names (against applicants and other manual names) and whitespace"""
        data = dict(self.base_data)
        # Applicant 1 is "અરજદાર એ (Applicant A)"
        app1_name = "અરજદાર એ (Applicant A)"
        data["LAND_RECORDS"] = [
            {
                "index": "1",
                "note_no": "500",
                "note_date": "01/01/2025",
                "note_details": "ડુપ્લિકેટ ટેસ્ટ",
                "owner_applicant_indices": [1, 2],
                "owner_other_names": [
                    "  ",                          # Empty/whitespace (should be ignored)
                    app1_name,                      # Exact duplicate of applicant 1 (should not be re-added)
                    "  નવો અન્ય કબજેદાર  ",        # Whitespace trimmed -> "નવો અન્ય કબજેદાર"
                    "નવો અન્ય કબજેદાર",            # Duplicate in other names (should not be duplicated)
                    ""                              # Empty (should be ignored)
                ]
            }
        ]
        out_docx = self.output_dir / "reg_scenario_11.docx"
        render_docx_template(str(self.template_file), data, str(out_docx))
        doc = Document(str(out_docx))
        owner_cell = doc.tables[1].rows[1].cells[5].text.strip()
        
        # Check normalized count of occurrences
        self.assertEqual(owner_cell.count("Applicant A"), 1)
        self.assertEqual(owner_cell.count("નવો અન્ય કબજેદાર"), 1)
        self.assertIn("Applicant B", owner_cell)
        self.assertNotIn("  ", owner_cell)

    def test_scenario_12_normalization_direct_unit_test(self):
        """Direct unit test of _normalize_context for LAND_RECORDS combining logic"""
        raw_data = {
            "APPLICANTS": [
                {"index": 1, "name": "A"},
                {"index": 2, "name": "B"},
                {"index": 3, "name": "C"},
                {"index": 4, "name": "D"}
            ],
            "LAND_RECORDS": [
                {
                    "note_no": "1",
                    "owner_applicant_indices": [1, 3],
                    "owner_other_names": ["Manual Person 1", "Manual Person 2"]
                }
            ]
        }
        normalized = _normalize_context(raw_data)
        land = normalized["LAND_RECORDS"][0]
        self.assertEqual(land["owner_names"], ["A", "C", "Manual Person 1", "Manual Person 2"])
        self.assertEqual(land["owner_name"], "A, C, Manual Person 1, Manual Person 2")

if __name__ == "__main__":
    unittest.main()
