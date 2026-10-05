import pytest
import os
import re
from docx import Document
from backend.services.docx_engine import (
    render_docx_template,
    extract_variables_from_docx,
)

def strip_markers(text: str) -> str:
    cleaned = re.sub(r'\[\[\[VAR_START:[^\]]+\]\]\]', '', text)
    cleaned = re.sub(r'\[\[\[VAR_MISSING:[^\]]+\]\]\]', '', cleaned)
    cleaned = re.sub(r'\[\[\[VAR_END\]\]\]', '', cleaned)
    return cleaned

def create_docx_with_paragraphs(paragraphs_list):
    os.makedirs("tests/temp", exist_ok=True)
    temp_path = os.path.join("tests/temp", f"sync_test_{os.urandom(4).hex()}.docx")
    doc = Document()
    for p in paragraphs_list:
        doc.add_paragraph(p)
    doc.save(temp_path)
    return temp_path

def get_temp_out():
    os.makedirs("tests/temp", exist_ok=True)
    return os.path.join("tests/temp", f"sync_out_{os.urandom(4).hex()}.docx")

def test_empty_vendors_zero_rows_all_occurrences():
    """Requirement 1: Zero vendor rows in every VENDORS occurrence, no Sample Name/Relation."""
    tpl_text = [
        "Header Notice",
        "{%p for vendor in VENDORS %}",
        "The Vendor: {{ vendor.name }} {{ vendor.relation }}",
        "{%p endfor %}",
        "Middle Body Clause",
        "{%p for vendor in VENDORS %}",
        "Signature: {{ vendor.name }} - {{ vendor.relation }}",
        "{%p endfor %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        render_docx_template(tmpl_path, {"VENDORS": []}, out_path, preview=True)
        doc = Document(out_path)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "The Vendor:" not in full_text
        assert "Signature:" not in full_text
        assert "Sample Name" not in full_text
        assert "Sample Relation" not in full_text
        assert "Header Notice" in full_text
        assert "Middle Body Clause" in full_text
    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)

def test_one_vendor_propagates_to_all_occurrences():
    """Requirement 2: One Vendor appears in all VENDORS occurrences identically."""
    tpl_text = [
        "{%p for vendor in VENDORS %}",
        "The Vendor: {{ vendor.name }} {{ vendor.relation }} aged {{ vendor.age }}",
        "{%p endfor %}",
        "Terms and Conditions",
        "{%p for vendor in VENDORS %}",
        "Sign: {{ vendor.name }} ({{ vendor.relation }})",
        "{%p endfor %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        data = {
            "VENDORS": [
                {"name": "Ramesh Patel", "relation": "S/o XYZ", "age": "45"}
            ]
        }
        render_docx_template(tmpl_path, data, out_path, preview=True)
        doc = Document(out_path)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        clean_text = strip_markers(full_text)
        assert "The Vendor: Ramesh Patel S/o XYZ aged 45" in clean_text
        assert "Sign: Ramesh Patel (S/o XYZ)" in clean_text
        assert clean_text.count("Ramesh Patel") == 2
        assert clean_text.count("S/o XYZ") == 2
        assert "Sample Name" not in clean_text
    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)

def test_two_vendors_ordering_across_all_occurrences():
    """Requirement 3: Two vendors appear in all occurrences with exact same ordering."""
    tpl_text = [
        "{%p for vendor in VENDORS %}",
        "Occ1: {{ vendor.index }}. {{ vendor.name }} - {{ vendor.relation }}",
        "{%p endfor %}",
        "{%p for vendor in VENDORS %}",
        "Occ2: {{ vendor.index }}. {{ vendor.name }} Signature",
        "{%p endfor %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        data = {
            "VENDORS": [
                {"name": "Ramesh Patel", "relation": "S/o XYZ"},
                {"name": "Suresh Patel", "relation": "S/o ABC"}
            ]
        }
        render_docx_template(tmpl_path, data, out_path, preview=True)
        doc = Document(out_path)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        clean_text = strip_markers(full_text)
        assert clean_text.count("Ramesh Patel") == 2
        assert clean_text.count("Suresh Patel") == 2
        # Verify ordering in Occ1
        occ1_ramesh = clean_text.find("Occ1: 1. Ramesh Patel")
        occ1_suresh = clean_text.find("Occ1: 2. Suresh Patel")
        assert 0 <= occ1_ramesh < occ1_suresh
        # Verify ordering in Occ2
        occ2_ramesh = clean_text.find("Occ2: 1. Ramesh Patel")
        occ2_suresh = clean_text.find("Occ2: 2. Suresh Patel")
        assert occ1_suresh < occ2_ramesh < occ2_suresh
    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)

def test_delete_one_vendor_updates_all_occurrences():
    """Requirement 4: Deleting vendor 2 immediately updates all occurrences."""
    tpl_text = [
        "{%p for vendor in VENDORS %}",
        "V: {{ vendor.name }}",
        "{%p endfor %}",
        "{%p for vendor in VENDORS %}",
        "Sign: {{ vendor.name }}",
        "{%p endfor %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        # After user deletes Suresh Patel:
        data = {
            "VENDORS": [
                {"name": "Ramesh Patel", "relation": "S/o XYZ"}
            ]
        }
        render_docx_template(tmpl_path, data, out_path, preview=True)
        doc = Document(out_path)
        clean_text = strip_markers("\n".join(p.text for p in doc.paragraphs))
        assert "Suresh Patel" not in clean_text
        assert clean_text.count("Ramesh Patel") == 2
    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)

def test_delete_final_vendor_clears_all_occurrences():
    """Requirement 5: Deleting final vendor leaves all occurrences empty without sample fallback."""
    tpl_text = [
        "{%p for vendor in VENDORS %}",
        "V: {{ vendor.name }}",
        "{%p endfor %}",
        "{%p for vendor in VENDORS %}",
        "Sign: {{ vendor.name }}",
        "{%p endfor %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        # All vendors deleted:
        data = {"VENDORS": []}
        render_docx_template(tmpl_path, data, out_path, preview=True)
        doc = Document(out_path)
        clean_text = strip_markers("\n".join(p.text for p in doc.paragraphs))
        assert "V:" not in clean_text
        assert "Sign:" not in clean_text
        assert "Sample Name" not in clean_text
        assert "Sample Relation" not in clean_text
    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)

def test_vendor_type_individual_vs_entity_visibility():
    """Requirement 6 & 7: VENDOR_TYPE INDIVIDUAL vs ENTITY conditional rendering."""
    tpl_text = [
        "{%p if VENDOR_TYPE == 'INDIVIDUAL' %}",
        "{%p for vendor in VENDORS %}",
        "Individual Vendor: {{ vendor.name }}",
        "{%p endfor %}",
        "{%p endif %}",
        "{%p if VENDOR_TYPE == 'ENTITY' %}",
        "Entity Name: {{ VENDOR_ENTITY_NAME }}",
        "{%p for rep in VENDOR_REPRESENTATIVES %}",
        "Representative: {{ rep.name }}",
        "{%p endfor %}",
        "{%p endif %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path1 = get_temp_out()
    out_path2 = get_temp_out()
    try:
        # Case 1: INDIVIDUAL
        data_ind = {
            "VENDOR_TYPE": "INDIVIDUAL",
            "VENDORS": [{"name": "Ramesh Patel"}],
            "VENDOR_ENTITY_NAME": "ABC Realty Pvt Ltd",
            "VENDOR_REPRESENTATIVES": [{"name": "Director Dinesh"}]
        }
        render_docx_template(tmpl_path, data_ind, out_path1, preview=True)
        text_ind = strip_markers("\n".join(p.text for p in Document(out_path1).paragraphs))
        assert "Individual Vendor: Ramesh Patel" in text_ind
        assert "Entity Name:" not in text_ind
        assert "Representative:" not in text_ind

        # Case 2: ENTITY
        data_ent = {
            "VENDOR_TYPE": "ENTITY",
            "VENDORS": [{"name": "Ramesh Patel"}],
            "VENDOR_ENTITY_NAME": "ABC Realty Pvt Ltd",
            "VENDOR_REPRESENTATIVES": [{"name": "Director Dinesh"}]
        }
        render_docx_template(tmpl_path, data_ent, out_path2, preview=True)
        text_ent = strip_markers("\n".join(p.text for p in Document(out_path2).paragraphs))
        assert "Individual Vendor:" not in text_ent
        assert "Entity Name: ABC Realty Pvt Ltd" in text_ent
        assert "Representative: Director Dinesh" in text_ent
    finally:
        for f in [tmpl_path, out_path1, out_path2]:
            if os.path.exists(f): os.remove(f)

def test_vendor_representatives_separate_collection():
    """Requirement 8: VENDOR_REPRESENTATIVES remains separate from VENDORS."""
    tpl_text = [
        "{%p for vendor in VENDORS %}",
        "Vendor: {{ vendor.name }}",
        "{%p endfor %}",
        "{%p for rep in VENDOR_REPRESENTATIVES %}",
        "Representative: {{ rep.name }}",
        "{%p endfor %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        data = {
            "VENDORS": [{"name": "Vendor A"}],
            "VENDOR_REPRESENTATIVES": [{"name": "Rep B"}]
        }
        render_docx_template(tmpl_path, data, out_path, preview=True)
        doc = Document(out_path)
        text = strip_markers("\n".join(p.text for p in doc.paragraphs))
        assert "Vendor: Vendor A" in text
        assert "Representative: Rep B" in text
        assert text.count("Vendor A") == 1
        assert text.count("Rep B") == 1
    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)

def test_purchasers_repeated_loop_occurrences():
    """Requirement 9: PURCHASERS repeated loop occurrences use one array."""
    tpl_text = [
        "{%p for p in PURCHASERS %}",
        "Purchaser Clause: {{ p.name }} ({{ p.pan }})",
        "{%p endfor %}",
        "Middle text",
        "{%p for p in PURCHASERS %}",
        "Purchaser Signature: {{ p.name }}",
        "{%p endfor %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        data = {
            "PURCHASERS": [
                {"name": "Anil Shah", "pan": "ABCDE1234F"},
                {"name": "Sunil Shah", "pan": "XYZPK5678M"}
            ]
        }
        render_docx_template(tmpl_path, data, out_path, preview=True)
        doc = Document(out_path)
        text = strip_markers("\n".join(p.text for p in doc.paragraphs))
        assert text.count("Anil Shah") == 2
        assert text.count("Sunil Shah") == 2
        assert text.count("ABCDE1234F") == 1
    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)

def test_plots_repeated_loop_occurrences():
    """Requirement 10: PLOTS repeated loop occurrences use one array."""
    tpl_text = [
        "{%p for plot in PLOTS %}",
        "Plot Schedule: No {{ plot.number }} Area {{ plot.area }}",
        "{%p endfor %}",
        "{%p for plot in PLOTS %}",
        "Boundary Plot: No {{ plot.number }}",
        "{%p endfor %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        data = {
            "PLOTS": [
                {"number": "101", "area": "500 sq.ft"},
                {"number": "102", "area": "600 sq.ft"}
            ]
        }
        render_docx_template(tmpl_path, data, out_path, preview=True)
        doc = Document(out_path)
        text = strip_markers("\n".join(p.text for p in doc.paragraphs))
        assert text.count("No 101") == 2
        assert text.count("No 102") == 2
    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)

def test_payments_repeated_loop_occurrences():
    """Requirement 11: PAYMENTS repeated loop occurrences use one array."""
    tpl_text = [
        "{%p for pay in PAYMENTS %}",
        "Payment Table: Rs. {{ pay.amount }} via {{ pay.mode }}",
        "{%p endfor %}",
        "{%p for pay in PAYMENTS %}",
        "Receipt: Rs. {{ pay.amount }}",
        "{%p endfor %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        data = {
            "PAYMENTS": [
                {"amount": "50,000", "mode": "Cheque"},
                {"amount": "25,000", "mode": "NEFT"}
            ]
        }
        render_docx_template(tmpl_path, data, out_path, preview=True)
        doc = Document(out_path)
        text = strip_markers("\n".join(p.text for p in doc.paragraphs))
        assert text.count("50,000") == 2
        assert text.count("25,000") == 2
    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)

def test_land_records_multi_applicant_mapping():
    """Requirement 12: Existing LAND_RECORDS multi-applicant occupant mapping remains intact."""
    tpl_text = [
        "{%p for lr in LAND_RECORDS %}",
        "Survey {{ lr.survey_no }} Owners: {{ lr.owner_name }}",
        "{%p endfor %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        data = {
            "APPLICANTS": [
                {"name": "Ghanshyam Patel"},
                {"name": "Vitthal Solanki"}
            ],
            "LAND_RECORDS": [
                {
                    "survey_no": "123",
                    "owner_applicant_indices": [1, 2],
                    "owner_other_names": ["External Co-owner"]
                }
            ]
        }
        render_docx_template(tmpl_path, data, out_path, preview=True)
        doc = Document(out_path)
        text = strip_markers("\n".join(p.text for p in doc.paragraphs))
        assert "Survey 123 Owners: Ghanshyam Patel, Vitthal Solanki, External Co-owner" in text
    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)

def test_entity_mode_representative_repeated_occurrences():
    """Requirement: VENDOR_TYPE == 'ENTITY' repeated loop occurrences render all representatives identically."""
    tpl_text = [
        "{%p if VENDOR_TYPE == 'ENTITY' %}",
        "Entity: {{ VENDOR_ENTITY_NAME }} ({{ VENDOR_ENTITY_PAN }})",
        "{%p for vendor in VENDOR_REPRESENTATIVES %}",
        "Rep Occ1: {{ vendor.index }}. {{ vendor.name }} {{ vendor.relation }} aged {{ vendor.age }}",
        "{%p endfor %}",
        "{%p endif %}",
        "Middle Terms Notice",
        "{%p if VENDOR_TYPE == 'ENTITY' %}",
        "{%p for vendor in VENDOR_REPRESENTATIVES %}",
        "Rep Occ2: {{ vendor.name }} - Signature",
        "{%p endfor %}",
        "{%p endif %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        # Case A: 1 representative
        data_1 = {
            "VENDOR_TYPE": "ENTITY",
            "VENDOR_ENTITY_NAME": "Alpha Corp Ltd",
            "VENDOR_ENTITY_PAN": "ABCDE1234F",
            "VENDOR_REPRESENTATIVES": [
                {"name": "Ramesh Patel", "relation": "S/o XYZ", "age": "45"}
            ]
        }
        render_docx_template(tmpl_path, data_1, out_path, preview=True)
        doc = Document(out_path)
        clean = strip_markers("\n".join(p.text for p in doc.paragraphs))
        assert "Rep Occ1: 1. Ramesh Patel S/o XYZ aged 45" in clean
        assert "Rep Occ2: Ramesh Patel - Signature" in clean
        assert clean.count("Ramesh Patel") == 2
        assert "Sample Name" not in clean

        # Case B: 2 representatives
        data_2 = {
            "VENDOR_TYPE": "ENTITY",
            "VENDOR_ENTITY_NAME": "Alpha Corp Ltd",
            "VENDOR_ENTITY_PAN": "ABCDE1234F",
            "VENDOR_REPRESENTATIVES": [
                {"name": "Ramesh Patel", "relation": "S/o XYZ", "age": "45"},
                {"name": "Suresh Patel", "relation": "S/o ABC", "age": "50"}
            ]
        }
        render_docx_template(tmpl_path, data_2, out_path, preview=True)
        clean2 = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert clean2.count("Ramesh Patel") == 2
        assert clean2.count("Suresh Patel") == 2
        assert "Rep Occ1: 1. Ramesh Patel" in clean2
        assert "Rep Occ1: 2. Suresh Patel" in clean2

        # Case C: Empty representatives -> 0 rows, no fallback
        data_empty = {
            "VENDOR_TYPE": "ENTITY",
            "VENDOR_ENTITY_NAME": "Alpha Corp Ltd",
            "VENDOR_ENTITY_PAN": "ABCDE1234F",
            "VENDOR_REPRESENTATIVES": []
        }
        render_docx_template(tmpl_path, data_empty, out_path, preview=True)
        clean_empty = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert "Rep Occ1:" not in clean_empty
        assert "Rep Occ2:" not in clean_empty
        assert "Sample Name" not in clean_empty
        assert "Sample Relation" not in clean_empty
        assert "Entity: Alpha Corp Ltd (ABCDE1234F)" in clean_empty
    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)

def test_vendor_mode_switching_and_live_preview():
    """Requirement: Switching INDIVIDUAL <-> ENTITY toggles visibility cleanly without sample fallback."""
    tpl_text = [
        "{%p if VENDOR_TYPE == 'INDIVIDUAL' %}",
        "{%p for vendor in VENDORS %}",
        "IndVendor: {{ vendor.name }}",
        "{%p endfor %}",
        "{%p endif %}",
        "{%p if VENDOR_TYPE == 'ENTITY' %}",
        "EntName: {{ VENDOR_ENTITY_NAME }}",
        "{%p for vendor in VENDOR_REPRESENTATIVES %}",
        "EntRep: {{ vendor.name }}",
        "{%p endfor %}",
        "{%p endif %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        # Step 1: INDIVIDUAL mode
        data_ind = {
            "VENDOR_TYPE": "INDIVIDUAL",
            "VENDORS": [{"name": "Ramesh Patel"}],
            "VENDOR_REPRESENTATIVES": [],
            "VENDOR_ENTITY_NAME": ""
        }
        render_docx_template(tmpl_path, data_ind, out_path, preview=True)
        txt1 = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert "IndVendor: Ramesh Patel" in txt1
        assert "EntName:" not in txt1
        assert "EntRep:" not in txt1

        # Step 2: Switch to ENTITY mode
        data_ent = {
            "VENDOR_TYPE": "ENTITY",
            "VENDORS": [{"name": "Ramesh Patel"}],
            "VENDOR_ENTITY_NAME": "Shree Ram Developers",
            "VENDOR_REPRESENTATIVES": [{"name": "Dinesh Shah"}]
        }
        render_docx_template(tmpl_path, data_ent, out_path, preview=True)
        txt2 = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert "IndVendor:" not in txt2
        assert "EntName: Shree Ram Developers" in txt2
        assert "EntRep: Dinesh Shah" in txt2

        # Step 3: Switch back to INDIVIDUAL mode
        data_ind_back = {
            "VENDOR_TYPE": "INDIVIDUAL",
            "VENDORS": [{"name": "Ramesh Patel"}],
            "VENDOR_ENTITY_NAME": "Shree Ram Developers",
            "VENDOR_REPRESENTATIVES": [{"name": "Dinesh Shah"}]
        }
        render_docx_template(tmpl_path, data_ind_back, out_path, preview=True)
        txt3 = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert "IndVendor: Ramesh Patel" in txt3
        assert "EntName:" not in txt3
        assert "EntRep:" not in txt3
    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)

def test_purchaser_entity_mode_representatives_full_lifecycle():
    """Requirement: PURCHASER_TYPE == 'ENTITY' representative multi-loop lifecycle and switching."""
    tpl_text = [
        "{%p if PURCHASER_TYPE == 'INDIVIDUAL' %}",
        "{%p for p in PURCHASERS %}",
        "IndPurchaser: {{ p.name }}",
        "{%p endfor %}",
        "{%p endif %}",
        "{%p if PURCHASER_TYPE == 'ENTITY' %}",
        "PurchaserEntity: {{ PURCHASER_ENTITY_NAME }}",
        "{%p for p in PURCHASER_REPRESENTATIVES %}",
        "Occ1 PurchaserRep: {{ p.index }}. {{ p.name }}",
        "{%p endfor %}",
        "Clause Separator",
        "{%p for p in PURCHASER_REPRESENTATIVES %}",
        "Occ2 PurchaserRep: {{ p.name }} Signature",
        "{%p endfor %}",
        "{%p endif %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        # Case A: INDIVIDUAL
        data_ind = {
            "PURCHASER_TYPE": "INDIVIDUAL",
            "PURCHASERS": [{"name": "Anil Shah"}],
            "PURCHASER_REPRESENTATIVES": [],
            "PURCHASER_ENTITY_NAME": ""
        }
        render_docx_template(tmpl_path, data_ind, out_path, preview=True)
        txt_ind = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert "IndPurchaser: Anil Shah" in txt_ind
        assert "PurchaserEntity:" not in txt_ind
        assert "PurchaserRep:" not in txt_ind

        # Case B: ENTITY with 2 representatives
        data_ent = {
            "PURCHASER_TYPE": "ENTITY",
            "PURCHASERS": [],
            "PURCHASER_ENTITY_NAME": "BlueStar Infra LLP",
            "PURCHASER_REPRESENTATIVES": [
                {"name": "Partner A"},
                {"name": "Partner B"}
            ]
        }
        render_docx_template(tmpl_path, data_ent, out_path, preview=True)
        txt_ent = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert "IndPurchaser:" not in txt_ent
        assert "PurchaserEntity: BlueStar Infra LLP" in txt_ent
        assert "Occ1 PurchaserRep: 1. Partner A" in txt_ent
        assert "Occ1 PurchaserRep: 2. Partner B" in txt_ent
        assert "Occ2 PurchaserRep: Partner A Signature" in txt_ent
        assert "Occ2 PurchaserRep: Partner B Signature" in txt_ent
        assert txt_ent.count("Partner A") == 2
        assert txt_ent.count("Partner B") == 2

        # Case C: Empty representatives
        data_empty = {
            "PURCHASER_TYPE": "ENTITY",
            "PURCHASERS": [],
            "PURCHASER_ENTITY_NAME": "BlueStar Infra LLP",
            "PURCHASER_REPRESENTATIVES": []
        }
        render_docx_template(tmpl_path, data_empty, out_path, preview=True)
        txt_empty = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert "PurchaserRep:" not in txt_empty
        assert "Sample Name" not in txt_empty
    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)

def test_unquoted_and_parenthesized_condition_extraction():
    """Requirement: extract_variables_from_docx supports unquoted and parenthesized conditions."""
    tpl_text = [
        "{%p if (VENDOR_TYPE == ENTITY) %}",
        "Name: {{ VENDOR_ENTITY_NAME }}",
        "{%p for rep in VENDOR_REPRESENTATIVES %}",
        "Rep: {{ rep.name }}",
        "{%p endfor %}",
        "{%p endif %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    try:
        res = extract_variables_from_docx(tmpl_path)
        conds = res.get("conditions", {})
        assert "VENDOR_ENTITY_NAME" in conds
        assert conds["VENDOR_ENTITY_NAME"]["field"] == "VENDOR_TYPE"
        assert conds["VENDOR_ENTITY_NAME"]["value"] == "ENTITY"
        assert "VENDOR_REPRESENTATIVES" in conds
        assert conds["VENDOR_REPRESENTATIVES"]["field"] == "VENDOR_TYPE"
        assert conds["VENDOR_REPRESENTATIVES"]["value"] == "ENTITY"
    finally:
        if os.path.exists(tmpl_path): os.remove(tmpl_path)

def test_conditional_live_preview_no_data_leak_between_individual_and_entity():
    """
    Requirement: Changing VENDOR_TYPE between INDIVIDUAL and ENTITY must NEVER leak
    stale data from the inactive branch into the rendered DOCX/Preview, even if
    the underlying form state preserves both collections and unbracketed loops exist.
    """
    tpl_text = [
        "{%p if VENDOR_TYPE == 'INDIVIDUAL' %}",
        "{%p for vendor in VENDORS %}",
        "Occ1 IndVendor: {{ vendor.index }}. {{ vendor.name }} {{ vendor.relation }}",
        "{%p endfor %}",
        "{%p endif %}",
        "{%p if VENDOR_TYPE == 'ENTITY' %}",
        "EntityName: {{ VENDOR_ENTITY_NAME }}",
        "EntityPAN: {{ VENDOR_ENTITY_PAN }}",
        "{%p for vendor in VENDOR_REPRESENTATIVES %}",
        "Occ1 Rep: {{ vendor.index }}. {{ vendor.name }} {{ vendor.relation }}",
        "{%p endfor %}",
        "{%p endif %}",
        "--- Signature Block ---",
        "{%p for vendor in VENDORS %}",
        "SignVendor: {{ vendor.name }}",
        "{%p endfor %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        # Step 1: User enters INDIVIDUAL vendor
        form_state = {
            "VENDOR_TYPE": "INDIVIDUAL",
            "VENDORS": [{"name": "ધર્મકુમાર શિવુભાઈ મકવાણા", "relation": "S/o XYZ"}],
            "VENDOR_ENTITY_NAME": "",
            "VENDOR_ENTITY_PAN": "",
            "VENDOR_REPRESENTATIVES": []
        }
        render_docx_template(tmpl_path, form_state, out_path, preview=True)
        txt1 = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert "Occ1 IndVendor: 1. ધર્મકુમાર શિવુભાઈ મકવાણા S/o XYZ" in txt1
        assert "SignVendor: ધર્મકુમાર શિવુભાઈ મકવાણા" in txt1
        assert "EntityName:" not in txt1
        assert "Occ1 Rep:" not in txt1

        # Step 2: User switches to ENTITY mode
        # Form state preserves VENDORS, but adds entity details and representative
        form_state["VENDOR_TYPE"] = "ENTITY"
        form_state["VENDOR_ENTITY_NAME"] = "ABC Pvt Ltd"
        form_state["VENDOR_ENTITY_PAN"] = "ABCDE1234F"
        form_state["VENDOR_REPRESENTATIVES"] = [
            {"name": "Ramesh Patel", "relation": "S/o XYZ", "age": "45"}
        ]
        render_docx_template(tmpl_path, form_state, out_path, preview=True)
        txt2 = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        # STALE INDIVIDUAL VENDOR MUST NOT LEAK
        assert "ધર્મકુમાર" not in txt2
        assert "Occ1 IndVendor:" not in txt2
        # ENTITY DATA MUST RENDER
        assert "EntityName: ABC Pvt Ltd" in txt2
        assert "EntityPAN: ABCDE1234F" in txt2
        assert "Occ1 Rep: 1. Ramesh Patel S/o XYZ" in txt2
        # SIGNATURE SECTION USES ACTIVE PARTY (REPRESENTATIVES)
        assert "SignVendor: Ramesh Patel" in txt2

        # Step 3: Add second representative
        form_state["VENDOR_REPRESENTATIVES"].append(
            {"name": "Suresh Patel", "relation": "S/o ABC", "age": "50"}
        )
        render_docx_template(tmpl_path, form_state, out_path, preview=True)
        txt3 = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert "ધર્મકુમાર" not in txt3
        assert "Occ1 Rep: 1. Ramesh Patel S/o XYZ" in txt3
        assert "Occ1 Rep: 2. Suresh Patel S/o ABC" in txt3
        assert "SignVendor: Ramesh Patel" in txt3
        assert "SignVendor: Suresh Patel" in txt3

        # Step 4: User switches back to INDIVIDUAL mode
        form_state["VENDOR_TYPE"] = "INDIVIDUAL"
        render_docx_template(tmpl_path, form_state, out_path, preview=True)
        txt4 = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        # PRESERVED INDIVIDUAL VENDOR RESTORED
        assert "Occ1 IndVendor: 1. ધર્મકુમાર શિવુભાઈ મકવાણા S/o XYZ" in txt4
        assert "SignVendor: ધર્મકુમાર શિવુભાઈ મકવાણા" in txt4
        # ENTITY DATA MUST BE EXCLUDED
        assert "EntityName:" not in txt4
        assert "ABC Pvt Ltd" not in txt4
        assert "Ramesh Patel" not in txt4
        assert "Suresh Patel" not in txt4
        assert "Occ1 Rep:" not in txt4

    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)


def test_conditional_purchaser_live_preview_no_data_leak():
    """Requirement: Symmetrical verification for PURCHASER_TYPE INDIVIDUAL <-> ENTITY."""
    tpl_text = [
        "{%p if PURCHASER_TYPE == 'INDIVIDUAL' %}",
        "{%p for p in PURCHASERS %}",
        "IndPurchaser: {{ p.name }}",
        "{%p endfor %}",
        "{%p endif %}",
        "{%p if PURCHASER_TYPE == 'ENTITY' %}",
        "PEntityName: {{ PURCHASER_ENTITY_NAME }}",
        "{%p for p in PURCHASER_REPRESENTATIVES %}",
        "Occ1 PRep: {{ p.index }}. {{ p.name }}",
        "{%p endfor %}",
        "{%p endif %}",
        "--- Purchaser Signature ---",
        "{%p for p in PURCHASERS %}",
        "SignPurchaser: {{ p.name }}",
        "{%p endfor %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        # Step 1: INDIVIDUAL
        data = {
            "PURCHASER_TYPE": "INDIVIDUAL",
            "PURCHASERS": [{"name": "Chhaganlal Patel"}],
            "PURCHASER_ENTITY_NAME": "",
            "PURCHASER_REPRESENTATIVES": []
        }
        render_docx_template(tmpl_path, data, out_path, preview=True)
        txt1 = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert "IndPurchaser: Chhaganlal Patel" in txt1
        assert "SignPurchaser: Chhaganlal Patel" in txt1
        assert "PEntityName:" not in txt1
        assert "Occ1 PRep:" not in txt1

        # Step 2: Switch to ENTITY with preserved PURCHASERS
        data["PURCHASER_TYPE"] = "ENTITY"
        data["PURCHASER_ENTITY_NAME"] = "Surat Realty LLP"
        data["PURCHASER_REPRESENTATIVES"] = [{"name": "Partner Vikram"}]
        render_docx_template(tmpl_path, data, out_path, preview=True)
        txt2 = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert "Chhaganlal Patel" not in txt2
        assert "IndPurchaser:" not in txt2
        assert "PEntityName: Surat Realty LLP" in txt2
        assert "Occ1 PRep: 1. Partner Vikram" in txt2
        assert "SignPurchaser: Partner Vikram" in txt2

        # Step 3: Switch back to INDIVIDUAL
        data["PURCHASER_TYPE"] = "INDIVIDUAL"
        render_docx_template(tmpl_path, data, out_path, preview=True)
        txt3 = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert "IndPurchaser: Chhaganlal Patel" in txt3
        assert "SignPurchaser: Chhaganlal Patel" in txt3
        assert "Surat Realty LLP" not in txt3
        assert "Partner Vikram" not in txt3
    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)


def test_entity_representative_signature_rendering_full_lifecycle():
    """
    Requirements:
    1. Individual Vendor signature uses VENDORS.
    2. Entity Vendor signature uses VENDOR_REPRESENTATIVES.
    3. Individual Purchaser signature uses PURCHASERS.
    4. Entity Purchaser signature uses PURCHASER_REPRESENTATIVES.
    5. Multiple Entity representatives appear in signature.
    6. Removing representative updates signature.
    7. Empty representatives produce zero signature rows.
    8. Old VENDORS data never appears in Entity signature.
    9. Switching Entity -> Individual restores Vendor signature.
    10. Switching Individual -> Entity changes signature source correctly.
    """
    tpl_text = [
        "{%p if VENDOR_TYPE == 'INDIVIDUAL' %}",
        "{%p for v in VENDORS %}",
        "MainIndVendor: {{ v.name }}",
        "{%p endfor %}",
        "{%p endif %}",
        "{%p if VENDOR_TYPE == 'ENTITY' %}",
        "VendorEntity: {{ VENDOR_ENTITY_NAME }}",
        "{%p for v in VENDOR_REPRESENTATIVES %}",
        "MainEntRep: {{ v.name }}",
        "{%p endfor %}",
        "{%p endif %}",
        "--- Vendor Signatures ---",
        "{%p for v in VENDORS %}",
        "VendorSignature: {{ v.index }}. {{ v.name }}",
        "{%p endfor %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        # 1. Individual mode
        state = {
            "VENDOR_TYPE": "INDIVIDUAL",
            "VENDORS": [
                {"name": "Dharmakumar"},
                {"name": "Ramesh Patel"}
            ],
            "VENDOR_ENTITY_NAME": "",
            "VENDOR_REPRESENTATIVES": []
        }
        render_docx_template(tmpl_path, state, out_path, preview=True)
        txt1 = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert "MainIndVendor: Dharmakumar" in txt1
        assert "MainIndVendor: Ramesh Patel" in txt1
        assert "VendorSignature: 1. Dharmakumar" in txt1
        assert "VendorSignature: 2. Ramesh Patel" in txt1
        assert "VendorEntity:" not in txt1
        assert "MainEntRep:" not in txt1

        # 2. Switch to Entity with 3 representatives
        state["VENDOR_TYPE"] = "ENTITY"
        state["VENDOR_ENTITY_NAME"] = "ABC Pvt Ltd"
        state["VENDOR_REPRESENTATIVES"] = [
            {"name": "Techno Gujarati"},
            {"name": "Suresh Patel"},
            {"name": "Mahesh Patel"}
        ]
        render_docx_template(tmpl_path, state, out_path, preview=True)
        txt2 = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert "Dharmakumar" not in txt2
        assert "MainIndVendor:" not in txt2
        assert "VendorEntity: ABC Pvt Ltd" in txt2
        assert "MainEntRep: Techno Gujarati" in txt2
        assert "MainEntRep: Suresh Patel" in txt2
        assert "MainEntRep: Mahesh Patel" in txt2
        # Signature block uses representatives
        assert "VendorSignature: 1. Techno Gujarati" in txt2
        assert "VendorSignature: 2. Suresh Patel" in txt2
        assert "VendorSignature: 3. Mahesh Patel" in txt2

        # 3. Delete representative 2 (Suresh Patel)
        state["VENDOR_REPRESENTATIVES"] = [
            {"name": "Techno Gujarati"},
            {"name": "Mahesh Patel"}
        ]
        render_docx_template(tmpl_path, state, out_path, preview=True)
        txt3 = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert "VendorSignature: 1. Techno Gujarati" in txt3
        assert "VendorSignature: 2. Mahesh Patel" in txt3
        assert "Suresh Patel" not in txt3

        # 4. Delete final representative -> 0 rows, no sample fallback
        state["VENDOR_REPRESENTATIVES"] = []
        render_docx_template(tmpl_path, state, out_path, preview=True)
        txt4 = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert "VendorSignature:" not in txt4
        assert "Sample Name" not in txt4
        assert "Sample Relation" not in txt4

        # 5. Switch back to INDIVIDUAL -> restores original vendor signatures
        state["VENDOR_TYPE"] = "INDIVIDUAL"
        render_docx_template(tmpl_path, state, out_path, preview=True)
        txt5 = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert "MainIndVendor: Dharmakumar" in txt5
        assert "MainIndVendor: Ramesh Patel" in txt5
        assert "VendorSignature: 1. Dharmakumar" in txt5
        assert "VendorSignature: 2. Ramesh Patel" in txt5
        assert "ABC Pvt Ltd" not in txt5
        assert "Techno Gujarati" not in txt5
        assert "Mahesh Patel" not in txt5

    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)


def test_entity_scalar_fields_never_corrupted_by_aliasing():
    """Verify entity scalar fields remain scalar, signature uses reps, and no raw python repr appears."""
    tpl_text = [
        "{%p if VENDOR_TYPE == 'INDIVIDUAL' %}",
        "{%p for vendor in VENDORS %}",
        "IndVendor: {{ vendor.name }}",
        "{%p endfor %}",
        "{%p endif %}",
        "{%p if VENDOR_TYPE == 'ENTITY' %}",
        "The Vendor Entity: {{ VENDOR_ENTITY_NAME }}",
        "Pancard No. {{ VENDOR_ENTITY_PAN }}",
        "{%p for vendor in VENDOR_REPRESENTATIVES %}",
        "Rep: {{ vendor.index }}. {{ vendor.name }}",
        "{%p endfor %}",
        "{%p endif %}",
        "--- Signature ---",
        "{%p for vendor in VENDORS %}",
        "Sign: {{ vendor.index }}. {{ vendor.name }}",
        "{%p endfor %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        # 1. ENTITY mode
        state = {
            "VENDOR_TYPE": "ENTITY",
            "VENDOR_ENTITY_NAME": "ABC Pvt Ltd",
            "VENDOR_ENTITY_PAN": "ABCDE1234F",
            "VENDOR_REPRESENTATIVES": [
                {"name": "Techno Gujarati"},
                {"name": "Suresh Patel"}
            ]
        }
        render_docx_template(tmpl_path, state, out_path, preview=True)
        full_text = "\n".join(p.text for p in Document(out_path).paragraphs)
        clean = strip_markers(full_text)

        # Entity scalar fields remain scalar strings
        assert "The Vendor Entity: ABC Pvt Ltd" in clean
        assert "Pancard No. ABCDE1234F" in clean
        # Representative loop renders
        assert "Rep: 1. Techno Gujarati" in clean
        assert "Rep: 2. Suresh Patel" in clean
        # Signature loop uses representatives
        assert "Sign: 1. Techno Gujarati" in clean
        assert "Sign: 2. Suresh Patel" in clean
        # No raw python list or dict representation anywhere
        assert "[{'index':" not in full_text
        assert "{'index':" not in full_text
        assert "IndVendor:" not in full_text

        # 2. INDIVIDUAL mode regression
        state_ind = {
            "VENDOR_TYPE": "INDIVIDUAL",
            "VENDOR_ENTITY_NAME": "ABC Pvt Ltd",
            "VENDOR_ENTITY_PAN": "ABCDE1234F",
            "VENDORS": [
                {"name": "Dharmakumar"},
                {"name": "Ramesh Patel"}
            ]
        }
        render_docx_template(tmpl_path, state_ind, out_path, preview=True)
        full_text_ind = "\n".join(p.text for p in Document(out_path).paragraphs)
        clean_ind = strip_markers(full_text_ind)

        assert "IndVendor: Dharmakumar" in clean_ind
        assert "IndVendor: Ramesh Patel" in clean_ind
        assert "Sign: 1. Dharmakumar" in clean_ind
        assert "Sign: 2. Ramesh Patel" in clean_ind
        assert "ABC Pvt Ltd" not in full_text_ind
        assert "ABCDE1234F" not in full_text_ind
        assert "[{'index':" not in full_text_ind
        assert "{'index':" not in full_text_ind
    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)


def test_purchaser_entity_scalar_fields_and_signature_aliasing():
    """Verify purchaser entity scalar fields remain scalar, signature uses reps, and no raw python repr appears."""
    tpl_text = [
        "{%p if PURCHASER_TYPE == 'INDIVIDUAL' %}",
        "{%p for purchaser in PURCHASERS %}",
        "IndPurchaser: {{ purchaser.name }}",
        "{%p endfor %}",
        "{%p endif %}",
        "{%p if PURCHASER_TYPE == 'ENTITY' %}",
        "Purchaser Entity: {{ PURCHASER_ENTITY_NAME }}",
        "Purchaser PAN: {{ PURCHASER_ENTITY_PAN }}",
        "{%p for purchaser in PURCHASER_REPRESENTATIVES %}",
        "PurchaserRep: {{ purchaser.index }}. {{ purchaser.name }}",
        "{%p endfor %}",
        "{%p endif %}",
        "--- Purchaser Signature ---",
        "{%p for purchaser in PURCHASERS %}",
        "PurchaserSign: {{ purchaser.index }}. {{ purchaser.name }}",
        "{%p endfor %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        # 1. ENTITY mode
        state = {
            "PURCHASER_TYPE": "ENTITY",
            "PURCHASER_ENTITY_NAME": "XYZ Infra Pvt Ltd",
            "PURCHASER_ENTITY_PAN": "XYZIN5678K",
            "PURCHASER_REPRESENTATIVES": [
                {"name": "Prakash Shah"},
                {"name": "Kiran Modi"}
            ]
        }
        render_docx_template(tmpl_path, state, out_path, preview=True)
        full_text = "\n".join(p.text for p in Document(out_path).paragraphs)
        clean = strip_markers(full_text)

        assert "Purchaser Entity: XYZ Infra Pvt Ltd" in clean
        assert "Purchaser PAN: XYZIN5678K" in clean
        assert "PurchaserRep: 1. Prakash Shah" in clean
        assert "PurchaserRep: 2. Kiran Modi" in clean
        assert "PurchaserSign: 1. Prakash Shah" in clean
        assert "PurchaserSign: 2. Kiran Modi" in clean
        assert "[{'index':" not in full_text
        assert "{'index':" not in full_text
        assert "IndPurchaser:" not in full_text

        # 2. INDIVIDUAL mode
        state_ind = {
            "PURCHASER_TYPE": "INDIVIDUAL",
            "PURCHASER_ENTITY_NAME": "XYZ Infra Pvt Ltd",
            "PURCHASER_ENTITY_PAN": "XYZIN5678K",
            "PURCHASERS": [
                {"name": "Vikram Rathod"}
            ]
        }
        render_docx_template(tmpl_path, state_ind, out_path, preview=True)
        full_text_ind = "\n".join(p.text for p in Document(out_path).paragraphs)
        clean_ind = strip_markers(full_text_ind)

        assert "IndPurchaser: Vikram Rathod" in clean_ind
        assert "PurchaserSign: 1. Vikram Rathod" in clean_ind
        assert "XYZ Infra Pvt Ltd" not in full_text_ind
        assert "XYZIN5678K" not in full_text_ind
        assert "[{'index':" not in full_text_ind
        assert "{'index':" not in full_text_ind
    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)


def test_docx_template_7d27c1de_exact_entity_rendering():
    """Verify exact template 7d27c1de renders entity scalar fields and signature with 0 raw python repr."""
    tpl_file = os.path.join("backend", "uploads", "templates_storage", "7d27c1de29a54f709d57b73cad1fc932.docx")
    if not os.path.exists(tpl_file):
        pytest.skip("Template 7d27c1de29a54f709d57b73cad1fc932.docx not present in local filesystem")

    out_path = get_temp_out()
    try:
        data = {
            "VENDOR_TYPE": "ENTITY",
            "VENDOR_ENTITY_NAME": "ABC Pvt Ltd",
            "VENDOR_ENTITY_PAN": "ABCDE1234F",
            "VENDOR_REPRESENTATIVES": [
                {"name": "Techno Gujarati", "relation": "S/o ABC", "age": "40"},
                {"name": "Suresh Patel", "relation": "S/o DEF", "age": "42"}
            ]
        }
        render_docx_template(tpl_file, data, out_path, preview=True)
        doc = Document(out_path)
        all_text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
        clean = strip_markers(all_text)

        assert "ABC Pvt Ltd" in clean
        assert "ABCDE1234F" in clean
        assert "Techno Gujarati" in clean
        assert "Suresh Patel" in clean
        assert "[{'index':" not in all_text
        assert "{'index':" not in all_text
    finally:
        if os.path.exists(out_path):
            os.remove(out_path)


def test_entity_representatives_no_duplication_in_main_and_signature_sections():
    """Verify Entity representatives are never duplicated in main section or signature section."""
    tpl_text = [
        "{%p if VENDOR_TYPE == 'INDIVIDUAL' %}",
        "{%p for vendor in VENDORS %}",
        "IndVendor: {{ vendor.name }}",
        "{%p endfor %}",
        "{%p endif %}",
        "{%p if VENDOR_TYPE == 'ENTITY' %}",
        "The Vendor Entity: {{ VENDOR_ENTITY_NAME }}",
        "Pancard No. {{ VENDOR_ENTITY_PAN }}",
        "{%p for vendor in VENDOR_REPRESENTATIVES %}",
        "Rep: {{ vendor.index }}. {{ vendor.name }}",
        "{%p endfor %}",
        "{%p endif %}",
        "--- Signature ---",
        "{%p for vendor in VENDORS %}",
        "Sign: {{ vendor.index }}. {{ vendor.name }}",
        "{%p endfor %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        # Case 1: Exactly 1 representative
        data_1 = {
            "VENDOR_TYPE": "ENTITY",
            "VENDOR_ENTITY_NAME": "Sagar LLP",
            "VENDOR_ENTITY_PAN": "ABCDE1234F",
            "VENDORS": [{"name": "Old Vendor Individual"}],
            "VENDOR_REPRESENTATIVES": [
                {"name": "Darshankumar", "relation": "son of Shivubhai", "pan": "ABCDE1234F"}
            ]
        }
        render_docx_template(tmpl_path, data_1, out_path, preview=True)
        txt1 = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert "The Vendor Entity: Sagar LLP" in txt1
        assert "Pancard No. ABCDE1234F" in txt1
        assert txt1.count("Rep: 1. Darshankumar") == 1
        assert txt1.count("Sign: 1. Darshankumar") == 1
        assert "Old Vendor Individual" not in txt1
        assert "IndVendor:" not in txt1

        # Case 2: Exactly 2 representatives
        data_2 = {
            "VENDOR_TYPE": "ENTITY",
            "VENDOR_ENTITY_NAME": "Sagar LLP",
            "VENDOR_ENTITY_PAN": "ABCDE1234F",
            "VENDORS": [{"name": "Old Vendor Individual"}],
            "VENDOR_REPRESENTATIVES": [
                {"name": "Darshankumar", "relation": "son of Shivubhai", "pan": "ABCDE1234F"},
                {"name": "Sagar", "relation": "son of Somabhai", "pan": "ABCDE1234F"}
            ]
        }
        render_docx_template(tmpl_path, data_2, out_path, preview=True)
        txt2 = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert txt2.count("Rep: 1. Darshankumar") == 1
        assert txt2.count("Rep: 2. Sagar") == 1
        assert txt2.count("Sign: 1. Darshankumar") == 1
        assert txt2.count("Sign: 2. Sagar") == 1
        assert "Old Vendor Individual" not in txt2

        # Case 3: Incoming payload with merged/duplicated representative array
        data_dup = {
            "VENDOR_TYPE": "ENTITY",
            "VENDOR_ENTITY_NAME": "Sagar LLP",
            "VENDOR_ENTITY_PAN": "ABCDE1234F",
            "VENDOR_REPRESENTATIVES": [
                {"name": "Darshankumar", "relation": "son of Shivubhai", "pan": "ABCDE1234F"},
                {"name": "Sagar", "relation": "son of Somabhai", "pan": "ABCDE1234F"},
                {"name": "Darshankumar", "relation": "son of Shivubhai", "pan": "ABCDE1234F"},
                {"name": "Sagar", "relation": "son of Somabhai", "pan": "ABCDE1234F"}
            ]
        }
        render_docx_template(tmpl_path, data_dup, out_path, preview=True)
        txt_dup = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert txt_dup.count("Rep: 1. Darshankumar") == 1
        assert txt_dup.count("Rep: 2. Sagar") == 1
        assert txt_dup.count("Sign: 1. Darshankumar") == 1
        assert txt_dup.count("Sign: 2. Sagar") == 1
        assert "Rep: 3." not in txt_dup
        assert "Rep: 4." not in txt_dup

        # Case 4: 3 representatives
        data_3 = {
            "VENDOR_TYPE": "ENTITY",
            "VENDOR_ENTITY_NAME": "Sagar LLP",
            "VENDOR_ENTITY_PAN": "ABCDE1234F",
            "VENDOR_REPRESENTATIVES": [
                {"name": "Darshankumar"},
                {"name": "Sagar"},
                {"name": "Pravin"}
            ]
        }
        render_docx_template(tmpl_path, data_3, out_path, preview=True)
        txt3 = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert txt3.count("Rep: 1. Darshankumar") == 1
        assert txt3.count("Rep: 2. Sagar") == 1
        assert txt3.count("Rep: 3. Pravin") == 1
        assert txt3.count("Sign: 1. Darshankumar") == 1
        assert txt3.count("Sign: 2. Sagar") == 1
        assert txt3.count("Sign: 3. Pravin") == 1
    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)


def test_purchaser_entity_representatives_no_duplication():
    """Verify Purchaser Entity representatives are never duplicated in main section or signature section."""
    tpl_text = [
        "{%p if PURCHASER_TYPE == 'INDIVIDUAL' %}",
        "{%p for p in PURCHASERS %}",
        "IndPurchaser: {{ p.name }}",
        "{%p endfor %}",
        "{%p endif %}",
        "{%p if PURCHASER_TYPE == 'ENTITY' %}",
        "Purchaser Entity: {{ PURCHASER_ENTITY_NAME }}",
        "Purchaser PAN: {{ PURCHASER_ENTITY_PAN }}",
        "{%p for p in PURCHASER_REPRESENTATIVES %}",
        "PRep: {{ p.index }}. {{ p.name }}",
        "{%p endfor %}",
        "{%p endif %}",
        "--- Purchaser Signature ---",
        "{%p for p in PURCHASERS %}",
        "PSign: {{ p.index }}. {{ p.name }}",
        "{%p endfor %}",
    ]
    tmpl_path = create_docx_with_paragraphs(tpl_text)
    out_path = get_temp_out()
    try:
        data_dup = {
            "PURCHASER_TYPE": "ENTITY",
            "PURCHASER_ENTITY_NAME": "Buildcon Ltd",
            "PURCHASER_ENTITY_PAN": "BLDCN1234E",
            "PURCHASERS": [{"name": "Old Purchaser Individual"}],
            "PURCHASER_REPRESENTATIVES": [
                {"name": "Kirit Shah", "pan": "BLDCN1234E"},
                {"name": "Anil Patel", "pan": "BLDCN1234E"},
                {"name": "Kirit Shah", "pan": "BLDCN1234E"},
                {"name": "Anil Patel", "pan": "BLDCN1234E"}
            ]
        }
        render_docx_template(tmpl_path, data_dup, out_path, preview=True)
        txt = strip_markers("\n".join(p.text for p in Document(out_path).paragraphs))
        assert "Purchaser Entity: Buildcon Ltd" in txt
        assert "Purchaser PAN: BLDCN1234E" in txt
        assert txt.count("PRep: 1. Kirit Shah") == 1
        assert txt.count("PRep: 2. Anil Patel") == 1
        assert txt.count("PSign: 1. Kirit Shah") == 1
        assert txt.count("PSign: 2. Anil Patel") == 1
        assert "PRep: 3." not in txt
        assert "Old Purchaser Individual" not in txt
        assert "IndPurchaser:" not in txt
    finally:
        for f in [tmpl_path, out_path]:
            if os.path.exists(f): os.remove(f)


def test_template_d2b220_exact_entity_representatives_rendering():
    """Verify real English plot template d2b22056 renders exactly once per representative in main cell and signature."""
    tpl_file = os.path.join("backend", "uploads", "templates_storage", "d2b22056060a49348c26ac5568e5a8b1.docx")
    if not os.path.exists(tpl_file):
        pytest.skip("Template d2b22056060a49348c26ac5568e5a8b1.docx not found")

    out_path = get_temp_out()
    try:
        data = {
            "VENDOR_TYPE": "ENTITY",
            "VENDOR_ENTITY_NAME": "Sagar LLP",
            "VENDOR_ENTITY_PAN": "ABCDE1234F",
            "VENDORS": [{"name": "Old Individual"}],
            "VENDOR_REPRESENTATIVES": [
                {"name": "Darshankumar", "relation": "son of Shivubhai", "pan": "ABCDE1234F", "aadhaar": "1234567890"},
                {"name": "Sagar", "relation": "son of Somabhai", "pan": "ABCDE1234F", "aadhaar": "1234567890"}
            ]
        }
        # 1. Preview mode
        render_docx_template(tpl_file, data, out_path, preview=True)
        doc = Document(out_path)
        main_vendor_cell_text = strip_markers(doc.tables[0].rows[0].cells[1].text)
        assert "Sagar LLP" in main_vendor_cell_text
        assert "ABCDE1234F" in main_vendor_cell_text
        assert main_vendor_cell_text.count("1. Darshankumar") == 1
        assert main_vendor_cell_text.count("2. Sagar") == 1
        assert "Old Individual" not in main_vendor_cell_text

        signature_text = strip_markers("\n".join(p.text for p in doc.paragraphs))
        assert signature_text.count("1. Darshankumar") == 1
        assert signature_text.count("2. Sagar") == 1
        assert "Old Individual" not in signature_text

        # 2. Final/PDF preview mode (preview=False)
        render_docx_template(tpl_file, data, out_path, preview=False)
        doc2 = Document(out_path)
        main_vendor_cell_text2 = doc2.tables[0].rows[0].cells[1].text
        assert "Sagar LLP" in main_vendor_cell_text2
        assert "ABCDE1234F" in main_vendor_cell_text2
        assert main_vendor_cell_text2.count("1. Darshankumar") == 1
        assert main_vendor_cell_text2.count("2. Sagar") == 1
        assert "Old Individual" not in main_vendor_cell_text2

        signature_text2 = "\n".join(p.text for p in doc2.paragraphs)
        assert signature_text2.count("1. Darshankumar") == 1
        assert signature_text2.count("2. Sagar") == 1
        assert "Old Individual" not in signature_text2
    finally:
        if os.path.exists(out_path):
            os.remove(out_path)


def test_vendor_and_purchaser_exact_switching_sequence():
    """
    Exact User Sequence Test:
    A. Select INDIVIDUAL -> enter individual vendor
    B. Switch to ENTITY -> enter Sagar LLP, ABCDE1234F, 2 reps (Darshankumar, Sagar)
    C. Verify no stale individual vendor in ENTITY mode
    D. Switch to INDIVIDUAL -> verify no entity name/PAN or reps
    E. Switch back to ENTITY -> verify no stale individual rows survive
    F. Symmetrical test for PURCHASER ENTITY mode
    """
    tpl_file = os.path.normpath(os.path.join(
        os.path.dirname(__file__), "..", "backend", "uploads", "templates_storage",
        "d2b22056060a49348c26ac5568e5a8b1.docx"
    ))
    if not os.path.exists(tpl_file):
        pytest.skip(f"Template 80 docx file not found at {tpl_file}")

    out_path = get_temp_out()
    try:
        # Step A & B: INDIVIDUAL mode
        data_step_a = {
            "VENDOR_TYPE": "INDIVIDUAL",
            "VENDORS": [{"name": "Dharmakumar Shivubhai Makwana"}],
            "PURCHASER_TYPE": "INDIVIDUAL",
            "PURCHASERS": [{"name": "Ghanshyam Patel"}]
        }
        render_docx_template(tpl_file, data_step_a, out_path, preview=True)
        doc_a = Document(out_path)
        text_a = strip_markers("\n".join(p.text for p in doc_a.paragraphs))
        cell_a = strip_markers(doc_a.tables[0].rows[0].cells[1].text)
        assert "Dharmakumar Shivubhai Makwana" in cell_a
        assert "Dharmakumar Shivubhai Makwana" in text_a
        assert "Sagar LLP" not in cell_a

        # Step C & D & E: Switch to ENTITY mode with stale VENDORS present in raw payload
        data_step_b = {
            "VENDOR_TYPE": "ENTITY",
            "VENDOR_ENTITY_NAME": "Sagar LLP",
            "VENDOR_ENTITY_PAN": "ABCDE1234F",
            "VENDOR_REPRESENTATIVES": [
                {"name": "Darshankumar"},
                {"name": "Sagar"}
            ],
            "VENDORS": [{"name": "Dharmakumar Shivubhai Makwana"}],
            "PURCHASER_TYPE": "ENTITY",
            "PURCHASER_ENTITY_NAME": "Purchaser Enterprises Ltd",
            "PURCHASER_ENTITY_PAN": "PPPPP9999P",
            "PURCHASER_REPRESENTATIVES": [{"name": "Kiran Patel"}],
            "PURCHASERS": [{"name": "Ghanshyam Patel"}]
        }
        render_docx_template(tpl_file, data_step_b, out_path, preview=True)
        doc_b = Document(out_path)
        vendor_cell_b = strip_markers(doc_b.tables[0].rows[0].cells[1].text)
        purchaser_cell_b = strip_markers(doc_b.tables[1].rows[0].cells[1].text)
        sig_text_b = strip_markers("\n".join(p.text for p in doc_b.paragraphs))

        # Main vendor section
        assert "Sagar LLP" in vendor_cell_b
        assert "ABCDE1234F" in vendor_cell_b
        assert "1. Darshankumar" in vendor_cell_b
        assert "2. Sagar" in vendor_cell_b
        assert "Dharmakumar" not in vendor_cell_b

        # Signature section
        assert "1. Darshankumar" in sig_text_b
        assert "2. Sagar" in sig_text_b
        assert "Dharmakumar" not in sig_text_b

        # Purchaser ENTITY section
        assert "Purchaser Enterprises Ltd" in purchaser_cell_b
        assert "PPPPP9999P" in purchaser_cell_b
        assert "Kiran Patel" in purchaser_cell_b
        assert "Ghanshyam Patel" not in purchaser_cell_b

        # Step F: Switch back to INDIVIDUAL with stale entity fields present
        data_step_f = {
            "VENDOR_TYPE": "INDIVIDUAL",
            "VENDORS": [{"name": "Dharmakumar Shivubhai Makwana"}],
            "VENDOR_ENTITY_NAME": "Sagar LLP",
            "VENDOR_ENTITY_PAN": "ABCDE1234F",
            "VENDOR_REPRESENTATIVES": [
                {"name": "Darshankumar"},
                {"name": "Sagar"}
            ],
            "PURCHASER_TYPE": "INDIVIDUAL",
            "PURCHASERS": [{"name": "Ghanshyam Patel"}],
            "PURCHASER_ENTITY_NAME": "Purchaser Enterprises Ltd",
            "PURCHASER_ENTITY_PAN": "PPPPP9999P",
            "PURCHASER_REPRESENTATIVES": [{"name": "Kiran Patel"}]
        }
        render_docx_template(tpl_file, data_step_f, out_path, preview=True)
        doc_f = Document(out_path)
        vendor_cell_f = strip_markers(doc_f.tables[0].rows[0].cells[1].text)
        purchaser_cell_f = strip_markers(doc_f.tables[1].rows[0].cells[1].text)
        sig_text_f = strip_markers("\n".join(p.text for p in doc_f.paragraphs))

        assert "Dharmakumar Shivubhai Makwana" in vendor_cell_f
        assert "Dharmakumar Shivubhai Makwana" in sig_text_f
        assert "Sagar LLP" not in vendor_cell_f
        assert "Darshankumar" not in vendor_cell_f
        assert "Ghanshyam Patel" in purchaser_cell_f
        assert "Purchaser Enterprises Ltd" not in purchaser_cell_f
    finally:
        if os.path.exists(out_path):
            os.remove(out_path)





